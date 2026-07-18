from fastapi import APIRouter, Depends, Query, HTTPException, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import and_, func
from typing import Optional, List
from datetime import datetime, date
from pydantic import BaseModel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO
from app.database import get_db
from app.models.models import AccessLog, User

router = APIRouter(prefix="/log-akses", tags=["Log Akses"])


def get_current_role(request: Request) -> str:
    """Helper untuk mengambil role dari token sesi."""
    from app.services.auth_service import decode_session_token
    token = request.cookies.get("session_token")
    if not token:
        raise HTTPException(status_code=401, detail="Sesi tidak valid")
    user = decode_session_token(token)
    return user["role"]


# ==========================================
# PYDANTIC SCHEMAS / MODELS
# ==========================================
class AccessLogDetail(BaseModel):
    id: int
    user_id: Optional[int] = None
    nama_user: Optional[str] = None
    nim_nip: Optional[str] = None
    ruangan_id: int
    waktu_akses: datetime
    metode: str
    foto_url: Optional[str] = None
    status: str
    keterangan: Optional[str] = None

    model_config = {"from_attributes": True}


class PaginationMeta(BaseModel):
    total: int
    page: int
    limit: int
    total_pages: int
    has_prev: bool
    has_next: bool


class PaginatedAccessLogResponse(BaseModel):
    data: List[AccessLogDetail]
    pagination: PaginationMeta


class BulkDeletePayload(BaseModel):
    ids: List[int]


# ==========================================
# ENDPOINTS
# ==========================================

@router.get("/", response_model=PaginatedAccessLogResponse)
def get_logs(
    ruangan_id:     Optional[int] = Query(None),
    status:         Optional[str] = Query(None),
    metode:         Optional[str] = Query(None),
    tanggal_dari:   Optional[str] = Query(None),
    tanggal_sampai: Optional[str] = Query(None),
    limit:          int = Query(10, ge=1, le=500),  # Default 10 agar mudah tes paginasi
    page:           int = Query(1, ge=1),
    db: Session = Depends(get_db)
):
    """Mengambil data log akses dengan filter dan paginasi."""
    q = db.query(AccessLog).options(joinedload(AccessLog.user))

    if ruangan_id:
        q = q.filter(AccessLog.ruangan_id == ruangan_id)
    if status:
        q = q.filter(AccessLog.status == status)
    if metode:
        q = q.filter(AccessLog.metode == metode)
    if tanggal_dari:
        q = q.filter(func.date(AccessLog.waktu_akses) >= tanggal_dari)
    if tanggal_sampai:
        q = q.filter(func.date(AccessLog.waktu_akses) <= tanggal_sampai)

    # Hitung total data sebelum di-slice oleh limit & offset
    total = q.count()
    offset = (page - 1) * limit
    
    # Ambil data sesuai halaman
    logs = q.order_by(AccessLog.waktu_akses.desc()).offset(offset).limit(limit).all()

    # Mapping data ke format schema Pydantic
    data_result = []
    for log in logs:
        data_result.append(AccessLogDetail(
            id          = log.id,
            user_id     = log.user_id,
            nama_user   = log.user.nama    if log.user else None,
            nim_nip     = log.user.nim_nip if log.user else None,
            ruangan_id  = log.ruangan_id,
            waktu_akses = log.waktu_akses,
            metode      = log.metode,
            foto_url    = log.foto_url,
            status      = log.status,
            keterangan  = log.keterangan,
        ))

    return {
        "data": data_result,
        "pagination": {
            "total":       total,
            "page":        page,
            "limit":       limit,
            "total_pages": (total + limit - 1) // limit,
            "has_prev":    page > 1,
            "has_next":    page * limit < total,
        }
    }


@router.delete("/bulk")
def hapus_bulk(
    payload: BulkDeletePayload,
    request: Request,
    db: Session = Depends(get_db)
):
    """Hapus beberapa log sekaligus berdasarkan list ID."""
    role = get_current_role(request)
    if role != "admin":
        raise HTTPException(status_code=403, detail="Hanya admin yang dapat menghapus log akses")

    if not payload.ids:
        return {"pesan": "Tidak ada ID yang dikirim", "dihapus": 0}

    deleted = db.query(AccessLog).filter(
        AccessLog.id.in_(payload.ids)
    ).delete(synchronize_session=False)
    db.commit()
    return {"pesan": f"{deleted} log berhasil dihapus", "dihapus": deleted}


@router.delete("/clear")
def hapus_semua(
    request: Request,
    ruangan_id: Optional[int] = Query(None),
    db: Session = Depends(get_db)
):
    """Hapus semua log. Jika ruangan_id diisi, hanya hapus log ruangan itu."""
    role = get_current_role(request)
    if role != "admin":
        raise HTTPException(status_code=403, detail="Hanya admin yang dapat menghapus seluruh log akses")

    q = db.query(AccessLog)
    if ruangan_id:
        q = q.filter(AccessLog.ruangan_id == ruangan_id)
    deleted = q.delete(synchronize_session=False)
    db.commit()
    label = f"ruangan {ruangan_id}" if ruangan_id else "semua ruangan"
    return {
        "pesan":   f"{deleted} log {label} berhasil dihapus",
        "dihapus": deleted
    }


@router.delete("/{log_id}")
def hapus_satu(log_id: int, request: Request, db: Session = Depends(get_db)):
    """Hapus satu log berdasarkan ID."""
    role = get_current_role(request)
    if role != "admin":
        raise HTTPException(status_code=403, detail="Hanya admin yang dapat menghapus log akses")

    log = db.query(AccessLog).filter(AccessLog.id == log_id).first()
    if not log:
        raise HTTPException(status_code=404, detail="Log tidak ditemukan")
    db.delete(log)
    db.commit()
    return {"pesan": "Log berhasil dihapus"}


@router.get("/export/excel")
def export_excel(
    ruangan_id:     Optional[int] = Query(None),
    status:         Optional[str] = Query(None),
    metode:         Optional[str] = Query(None),
    tanggal_dari:   Optional[str] = Query(None),
    tanggal_sampai: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    q = db.query(AccessLog).options(joinedload(AccessLog.user))
    if ruangan_id:
        q = q.filter(AccessLog.ruangan_id == ruangan_id)
    if status:
        q = q.filter(AccessLog.status == status)
    if metode:
        q = q.filter(AccessLog.metode == metode)
    if tanggal_dari:
        q = q.filter(func.date(AccessLog.waktu_akses) >= tanggal_dari)
    if tanggal_sampai:
        q = q.filter(func.date(AccessLog.waktu_akses) <= tanggal_sampai)

    logs = q.order_by(AccessLog.waktu_akses.desc()).limit(5000).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Log Akses"

    hdr_font  = Font(bold=True, color="FFFFFF", size=11)
    hdr_fill  = PatternFill("solid", fgColor="1F3864")
    hdr_align = Alignment(horizontal="center", vertical="center")

    headers    = ["No", "Waktu Akses", "Nama Pengguna", "NIM/NIP", "Metode", "Status", "Keterangan"]
    col_widths = [5, 22, 25, 20, 15, 12, 30]

    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 22

    for i, log in enumerate(logs, 1):
        waktu = log.waktu_akses.strftime("%d/%m/%Y %H:%M:%S") if log.waktu_akses else ""
        row = [
            i, waktu,
            log.user.nama    if log.user else "Tidak dikenal",
            log.user.nim_nip if log.user else "-",
            log.metode, log.status, log.keterangan or ""
        ]
        for col, val in enumerate(row, 1):
            cell = ws.cell(row=i+1, column=col, value=val)
            cell.alignment = Alignment(vertical="center")
            if i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="EBF3FB")

    ws.insert_rows(1)
    ws.merge_cells("A1:G1")
    tc = ws["A1"]
    tc.value = (f"Log Akses — Teknik Komputer Polsri "
                f"(Export: {datetime.now().strftime('%d/%m/%Y %H:%M')})")
    tc.font  = Font(bold=True, size=12, color="1F3864")
    tc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 20

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"log_akses_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"}
    )